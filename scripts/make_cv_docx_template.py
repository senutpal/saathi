"""One-shot generator for templates_docx/cv_template.docx (P04-S7 source artifact).

Run from repo root:
    uv run python scripts/make_cv_docx_template.py
The generated .docx is committed; rerun only when the CV docx design changes.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1] / "saathi" / "templates_docx" / "cv_template.docx"
)


def _heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(13)


def _entry_table(document: Document, loop_var: str, columns: list[str]) -> None:
    table = document.add_table(rows=4, cols=len(columns))
    table.style = "Table Grid"
    for index, label in enumerate(columns):
        cell_paragraph = table.rows[0].cells[index].paragraphs[0]
        run = cell_paragraph.add_run(label)
        run.bold = True
    table.rows[1].cells[0].paragraphs[0].add_run(f"{{%tr for {loop_var} in {loop_var} %}}")
    for index, field in enumerate(columns):
        table.rows[2].cells[index].paragraphs[0].add_run(f"{{{{ {loop_var}.{field} }}}}")
    table.rows[3].cells[0].paragraphs[0].add_run("{%tr endfor %}")


def build_template() -> None:
    document = Document()

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run("{{ name }}")
    run.bold = True
    run.font.size = Pt(24)

    headline = document.add_paragraph()
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headline_run = headline.add_run("{{ headline }}")
    headline_run.italic = True
    headline_run.font.size = Pt(12)

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact.add_run("{{ contact }}")
    contact_run.font.size = Pt(10)
    contact_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    _heading(document, "Summary")
    document.add_paragraph("{{ summary }}")

    _heading(document, "Skills")
    document.add_paragraph("{{ skills_text }}")

    _heading(document, "Education")
    _entry_table(document, "education", ["title", "org", "span"])

    document.add_paragraph()
    _heading(document, "Experience")
    _entry_table(document, "work", ["role", "org", "span"])

    document.add_paragraph()
    _heading(document, "Interests")
    document.add_paragraph("Likes: {{ likes_text }}")
    document.add_paragraph("Dislikes: {{ dislikes_text }}")

    footer = document.add_paragraph()
    footer_run = footer.add_run("Generated {{ generated_at_text }} via Saathi")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(TEMPLATE_PATH)
    print(f"wrote {TEMPLATE_PATH}")


if __name__ == "__main__":
    build_template()
